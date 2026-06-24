import argparse
import json
import os
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from go_pipeline.scripts.helper.pipeline_state import PipelineState

def load_json(path):
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def find_signatures(input_path):
    """Find all non-diluted signature folders."""
    signatures = []
    for entry in Path(input_path).iterdir():
        if entry.is_dir() and not entry.name.startswith("diluted_"):
            signatures.append(entry)
    return signatures


def load_manifold(path):
    data = load_json(path)
    go_map = {}

    for group in data.get("groups", []):
        for term in group.get("terms", []):
            go_id = term.get("go_id")

            if not go_id:
                continue

            go_map[go_id] = {
                "definition": term.get("definition", ""),
                "gene_symbols": term.get("gene_symbols", {})
            }

    return go_map


def filter_terms(ranking, apply_zero_filter=True):
    passing = []
    for term in ranking:
        robustness_values = term["robustness_score"]
        original = robustness_values.get("Original", 0)

        if original < 30:
            continue

        if apply_zero_filter:
            step_values = [v for k, v in robustness_values.items() if k.startswith("Step_")]
            zero_count = sum(1 for v in step_values if v == 0.0)
            if zero_count >= 2:
                continue

        passing.append(term)
    return passing


def build_signature_dict(terms, hierarchy_map, manifold_map, signature_name, ontology):
    result = {}
    for term in terms:
        go_id = term["go_id"]
        go_name = term["go_name"]
        initial_robustness = term["robustness_score"]["Original"]
        mean_robustness = term["score_components"]["mean_robustness"]

        manifold_entry = manifold_map.get(go_id, {})
        definition = manifold_entry.get("definition", "")
        gene_symbols_map = manifold_entry.get("gene_symbols", {})

        # Hierarchy map is only available for BP; skip gene lookup for MF/CC
        if ontology == "BP" and hierarchy_map is not None:
            if go_id not in hierarchy_map:
                print(f"  WARNING [{signature_name}]: GO term {go_id} not found in hierarchy_map_{ontology.lower()}.json")
                genes = {}
            else:
                genes = {}
                for gene in hierarchy_map[go_id]:
                    if gene not in gene_symbols_map:
                        print(f"  WARNING [{signature_name}]: Gene '{gene}' not found in manifold gene_symbols for {go_id}")
                        gene_def = ""
                    else:
                        gene_def = gene_symbols_map[gene]
                    genes[gene] = {"definition": gene_def}
        else:
            # For MF and CC: use gene_symbols directly from manifold
            genes = {gene: {"definition": gene_def} for gene, gene_def in gene_symbols_map.items()}

        result[go_id] = {
            "go_name": go_name,
            "initial_robustness": initial_robustness,
            "mean_robustness": mean_robustness,
            "definition": definition,
            "genes": genes,
        }

    return result


def save_json(data, output_path):
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)


def save_excel(sig_dict, output_path, signature_name):
    wb = Workbook()
    ws = wb.active
    ws.title = signature_name[:31]

    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", start_color="2F4F8F")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    go_fill_even = PatternFill("solid", start_color="DCE6F1")
    go_fill_odd = PatternFill("solid", start_color="EBF2FA")
    gene_fill_even = PatternFill("solid", start_color="F5F5F5")
    gene_fill_odd = PatternFill("solid", start_color="FFFFFF")

    go_font = Font(name="Arial", bold=True, size=9)
    gene_font = Font(name="Arial", size=9)
    wrap_align = Alignment(vertical="top", wrap_text=True)
    center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

    thin = Side(style="thin", color="AAAAAA")
    medium = Side(style="medium", color="2F4F8F")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    medium_border = Border(left=medium, right=medium, top=medium, bottom=medium)

    headers = [
        "GO ID", "GO Name", "Initial Robustness (%)", "Mean Robustness (%)",
        "GO Definition", "Gene Symbol", "Gene Definition"
    ]
    col_widths = [14, 35, 16, 16, 45, 14, 60]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = medium_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 30

    current_row = 2
    go_block_index = 0

    for go_id, info in sig_dict.items():
        genes = info["genes"]
        n_genes = len(genes) if genes else 1

        is_even = go_block_index % 2 == 0
        go_fill = go_fill_even if is_even else go_fill_odd
        gene_fill = gene_fill_even if is_even else gene_fill_odd

        go_data = [
            go_id,
            info["go_name"],
            round(info["initial_robustness"], 2),
            round(info["mean_robustness"], 2),
            info["definition"],
        ]

        start_row = current_row

        if genes:
            for gene_symbol, gene_info in genes.items():

                for col_idx, value in enumerate(go_data, start=1):
                    cell = ws.cell(
                        row=current_row,
                        column=col_idx,
                        value=value if current_row == start_row else None
                    )
                    cell.fill = go_fill
                    cell.font = go_font
                    cell.alignment = center_align if col_idx in (3, 4) else wrap_align
                    cell.border = thin_border

                ws.cell(row=current_row, column=6, value=gene_symbol).font = gene_font
                ws.cell(row=current_row, column=6).fill = gene_fill
                ws.cell(row=current_row, column=6).alignment = wrap_align
                ws.cell(row=current_row, column=6).border = thin_border

                ws.cell(row=current_row, column=7, value=gene_info["definition"]).font = gene_font
                ws.cell(row=current_row, column=7).fill = gene_fill
                ws.cell(row=current_row, column=7).alignment = wrap_align
                ws.cell(row=current_row, column=7).border = thin_border

                current_row += 1
        else:
            for col_idx, value in enumerate(go_data, start=1):
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                cell.fill = go_fill
                cell.font = go_font
                cell.alignment = center_align if col_idx in (3, 4) else wrap_align
                cell.border = thin_border

            ws.cell(row=current_row, column=6, value="")
            ws.cell(row=current_row, column=7, value="")
            current_row += 1

        if n_genes > 1:
            end_row = start_row + n_genes - 1
            for col_idx in range(1, 6):
                ws.merge_cells(start_row=start_row, start_column=col_idx,
                               end_row=end_row, end_column=col_idx)

        go_block_index += 1

    ws.freeze_panes = "A2"
    wb.save(output_path)


def process_ontology(sig_dir, sig_name, ontology, mode):
    """Process a single ontology for a given signature."""

    ont_lower = ontology.lower()

    # New path structure: Dilutions/{ontology}/{mode}/
    ranking_path = sig_dir / "Dilutions" / ontology / mode / f"{sig_name}_ranking.json"
    manifold_path = sig_dir / "parameter_analysis" / ontology / f"manifold_analysis_{ontology}.json"
    output_dir = sig_dir / "Dilutions" / ontology / mode

    # Hierarchy map only exists for BP
    if ontology == "BP":
        hierarchy_path = sig_dir / "representatives_analysis" / "hierarchy_map_bp.json"
        if not hierarchy_path.exists():
            print(f"  ERROR [{ontology}]: Hierarchy map not found: {hierarchy_path}")
            return
        hierarchy_map = load_json(hierarchy_path)
    else:
        hierarchy_map = None

    if not ranking_path.exists():
        print(f"  ERROR [{ontology}]: Ranking file not found: {ranking_path}")
        return
    if not manifold_path.exists():
        print(f"  ERROR [{ontology}]: Manifold file not found: {manifold_path}")
        return

    ranking_data = load_json(ranking_path)
    manifold_map = load_manifold(manifold_path)
    ranking = ranking_data.get("ranking", [])

    # Filtered (both filters)
    terms_both = filter_terms(ranking, apply_zero_filter=True)
    sig_dict_both = build_signature_dict(terms_both, hierarchy_map, manifold_map, sig_name, ontology)
    save_json(sig_dict_both, output_dir / f"{sig_name}_{ont_lower}_robustness_analysis_filtered.json")
    save_excel(sig_dict_both, output_dir / f"{sig_name}_{ont_lower}_robustness_analysis_filtered.xlsx", sig_name)

    # Robustness-only filter
    terms_robustness_only = filter_terms(ranking, apply_zero_filter=False)
    sig_dict_robustness = build_signature_dict(terms_robustness_only, hierarchy_map, manifold_map, sig_name, ontology)
    save_json(sig_dict_robustness, output_dir / f"{sig_name}_{ont_lower}_robustness_analysis.json")
    save_excel(sig_dict_robustness, output_dir / f"{sig_name}_{ont_lower}_robustness_analysis.xlsx", sig_name)

    print(f"  [{ontology}] Done — {len(terms_both)} filtered terms, {len(terms_robustness_only)} robustness-only terms")


def main():
    parser = argparse.ArgumentParser(description="Analyze GO term robustness per signature.")
    parser.add_argument("--input_path", required=True)
    parser.add_argument("--mode", required=True, choices=['cumulative', 'fixed'],
                        help="Dilution mode used in dilute_analysis.py")
    parser.add_argument("--ontology", default='BP', choices=['BP', 'MF', 'CC', 'all'],
                        help="GO ontology to analyze: 'BP', 'MF', 'CC', or 'all' (default: BP)")
    parser.add_argument("--state_file", default=None,
                        help="Path to pipeline status data")
    args = parser.parse_args()

    input_path = Path(args.input_path)
    ontologies = ['BP', 'MF', 'CC'] if args.ontology == 'all' else [args.ontology]
    signatures = find_signatures(input_path)

    if not signatures:
        print("No non-diluted signature folders found.")
        return

    state = PipelineState(args.state_file or os.path.join(str(input_path), ".pipeline_state.json"))

    print(f"Found {len(signatures)} signature(s): {[s.name for s in signatures]}")
    print(f"Mode: {args.mode.upper()} | Ontologies: {', '.join(ontologies)}")

    for sig_dir in signatures:
        sig_name = sig_dir.name
        print(f"\nProcessing: {sig_name}")

        for ontology in ontologies:
            work_key = f"{sig_name}::{ontology}::{args.mode}"

            if state.is_done("create_summary_data", work_key):
                continue

            try:
                process_ontology(sig_dir, sig_name, ontology, args.mode)
                ont_lower = ontology.lower()
                out_dir = sig_dir / "Dilutions" / ontology / args.mode
                expected_files = [
                    out_dir / f"{sig_name}_{ont_lower}_robustness_analysis_filtered.json",
                    out_dir / f"{sig_name}_{ont_lower}_robustness_analysis.json",
                ]
                if all(f.exists() for f in expected_files):
                    state.mark_done("create_summary_data", work_key)
                else:
                    state.mark_failed(
                        "create_summary_data", work_key,
                        "process_ontology produced no output files (see ERROR output above, e.g. missing ranking/manifold/hierarchy file)"
                    )
            except Exception as e:
                state.mark_failed("create_summary_data", work_key, str(e))
                print(f"  Error in {sig_name}/{ontology}: {e} -> continuing with next combination")
                continue


if __name__ == "__main__":
    main()
